"""
OCR Text Editor backend for S2POT.

Loads the OCR text of a single page from the generated result files and, after a manual
correction by the user, propagates the corrected text to every affected artefact:

  * searchable PDFs (invisible text layer rebuilt, page image kept byte-for-byte),
  * Word (.docx) documents (page paragraphs replaced, illustrations/tables kept),
  * plain text (.txt), hOCR (.hocr),
  * word coordinates (_words.json / .csv / .xlsx),
  * s2pot_metadata.json (chars, confidence and file aggregates).
"""
import os
import re
import csv
import json
import copy
import html
import difflib
import statistics
from datetime import datetime

import pymupdf as fitz

from core.metadata_config import load_folder_metadata, save_folder_metadata, estimate_text_layer_confidence
from core.pdf_builder import insert_invisible_text_items
from core.data_exporter import write_words_json, write_words_csv, write_words_xlsx, WORDS_KEYS

try:
    import docx
    from docx.shared import Pt
    from docx.oxml.ns import qn
except Exception:  # pragma: no cover
    docx = None

MANUAL_EDIT_CONFIDENCE = 100.0


# ----------------------------------------------------------------------------------------
# Result file discovery
# ----------------------------------------------------------------------------------------
def _norm(p):
    return os.path.normcase(os.path.abspath(p)) if p else ""


def _first_existing(paths):
    seen = set()
    out = []
    for p in paths:
        if p and os.path.isfile(p) and _norm(p) not in seen:
            seen.add(_norm(p))
            out.append(p)
    return out


def find_result_files(source_path, output_dir=None, data_dir=None):
    """Returns a dict with every generated result file that belongs to the source document."""
    base = os.path.splitext(os.path.basename(source_path))[0]
    out = output_dir or ""
    data = data_dir or out
    src_norm = _norm(source_path)

    def j(*parts):
        return os.path.join(*parts) if parts[0] else ""

    pdf_cands = [
        j(out, "PDF-Compressed", f"{base}_compressed.pdf"),
        j(out, f"{base}_compressed.pdf"),
        j(out, "PDF-OriginalImages", f"{base}_original.pdf"),
        j(out, "PDF-OriginalImages", f"{base}.pdf"),
        j(out, f"{base}_original.pdf"),
        j(out, f"{base}.pdf"),
    ]
    pdfs = [p for p in _first_existing(pdf_cands) if _norm(p) != src_norm]

    return {
        "base_name": base,
        "pdf": pdfs,
        "docx": _first_existing([j(out, "DOCX", f"{base}.docx"), j(out, f"{base}.docx")]),
        "txt": [p for p in _first_existing([j(data, "Text", f"{base}.txt"), j(data, f"{base}.txt")]) if _norm(p) != src_norm],
        "hocr": _first_existing([j(data, "hOCR", f"{base}.hocr"), j(data, f"{base}.hocr")]),
        "words_json": _first_existing([j(data, "Word_Coordinates", f"{base}_words.json"), j(data, f"{base}_words.json")]),
        "words_csv": _first_existing([j(data, "Word_Coordinates", f"{base}_words.csv"), j(data, f"{base}_words.csv")]),
        "words_xlsx": _first_existing([j(data, "Word_Coordinates", f"{base}_words.xlsx"), j(data, f"{base}_words.xlsx")]),
    }


# ----------------------------------------------------------------------------------------
# Line grouping helpers
# ----------------------------------------------------------------------------------------
def _group_words_into_lines(words):
    """
    words: list of dicts {"text", "bbox": (x0, y0, x1, y1), ...}
    Returns list of lines, each {"bbox": (x0, y0, x1, y1), "words": [...], "text": str}
    """
    words = [w for w in words if str(w.get("text", "")).strip()]
    if not words:
        return []
    heights = [max(1e-6, w["bbox"][3] - w["bbox"][1]) for w in words]
    thr = max(1e-6, statistics.median(heights) * 0.55)

    sorted_w = sorted(words, key=lambda w: ((w["bbox"][1] + w["bbox"][3]) / 2.0, w["bbox"][0]))
    lines = []
    cur = []
    cur_mid = None
    for w in sorted_w:
        mid = (w["bbox"][1] + w["bbox"][3]) / 2.0
        if cur_mid is None or abs(mid - cur_mid) <= thr:
            cur.append(w)
            cur_mid = sum((b["bbox"][1] + b["bbox"][3]) / 2.0 for b in cur) / float(len(cur))
        else:
            lines.append(cur)
            cur = [w]
            cur_mid = mid
    if cur:
        lines.append(cur)

    result = []
    for ln in lines:
        ln = sorted(ln, key=lambda b: b["bbox"][0])
        x0 = min(b["bbox"][0] for b in ln)
        y0 = min(b["bbox"][1] for b in ln)
        x1 = max(b["bbox"][2] for b in ln)
        y1 = max(b["bbox"][3] for b in ln)
        result.append({
            "bbox": (x0, y0, x1, y1),
            "words": ln,
            "text": " ".join(str(b["text"]).strip() for b in ln)
        })
    return result


def _count_chars(text):
    return sum(len(t) for t in (text or "").split())


# ----------------------------------------------------------------------------------------
# Reading page OCR text
# ----------------------------------------------------------------------------------------
def _pdf_page_words(pdf_path, page_idx):
    """Returns (words_in_points, page_w_pt, page_h_pt, img_w_px, img_h_px) or None."""
    try:
        with fitz.open(pdf_path) as doc:
            if page_idx >= len(doc):
                return None
            page = doc[page_idx]
            words = [{"text": w[4], "bbox": (w[0], w[1], w[2], w[3])} for w in page.get_text("words")]
            img_w = img_h = None
            try:
                infos = page.get_image_info()
                if infos:
                    big = max(infos, key=lambda i: (i.get("width", 0) or 0) * (i.get("height", 0) or 0))
                    img_w, img_h = big.get("width"), big.get("height")
            except Exception:
                pass
            return words, page.rect.width, page.rect.height, img_w, img_h
    except Exception:
        return None


def _load_words_list(files):
    """Loads the full word-coordinates list from JSON (preferred), CSV or XLSX."""
    if files.get("words_json"):
        try:
            with open(files["words_json"][0], "r", encoding="utf-8") as f:
                payload = json.load(f)
            return payload.get("words", []), payload
        except Exception:
            pass
    if files.get("words_csv"):
        try:
            rows = []
            with open(files["words_csv"][0], "r", encoding="utf-8-sig", newline="") as f:
                reader = csv.reader(f)
                next(reader, None)
                for r in reader:
                    if len(r) >= len(WORDS_KEYS):
                        rows.append(_coerce_word_row(r))
            return rows, None
        except Exception:
            pass
    if files.get("words_xlsx"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(files["words_xlsx"][0], read_only=True)
            ws = wb.active
            rows = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i == 0 or r is None or len(r) < len(WORDS_KEYS):
                    continue
                rows.append(_coerce_word_row(list(r)))
            wb.close()
            return rows, None
        except Exception:
            pass
    return None, None


def _coerce_word_row(r):
    d = dict(zip(WORDS_KEYS, r))
    for k in ("page", "word_num"):
        try:
            d[k] = int(float(d[k]))
        except Exception:
            d[k] = 0
    for k in ("x0", "y0", "x1", "y1", "width", "height", "font_size_pt"):
        try:
            v = float(d[k])
            d[k] = int(v) if v.is_integer() else v
        except Exception:
            d[k] = 0
    d["text"] = "" if d.get("text") is None else str(d["text"])
    return d


def _parse_txt_sections(content):
    """Parses '--- Page N ---' sections. Returns (preamble, ordered list of [num, text])."""
    markers = list(re.finditer(r"--- Page (\d+) ---\n?", content))
    if not markers:
        return content, []
    preamble = content[:markers[0].start()]
    sections = []
    for i, m in enumerate(markers):
        end = markers[i + 1].start() if i + 1 < len(markers) else len(content)
        body = content[m.end():end]
        if i + 1 < len(markers):
            body = body[:-2] if body.endswith("\n\n") else body.rstrip("\n")
        sections.append([int(m.group(1)), body])
    return preamble, sections


def load_page_ocr_text(source_path, page_idx, output_dir=None, data_dir=None):
    """
    Returns {"text", "chars", "confidence", "source", "has_ocr", "files"} for the requested page.
    Line structure is derived from word bounding boxes, so line i of the returned text maps
    onto line box i of the page when the corrected text is saved back.
    """
    files = find_result_files(source_path, output_dir, data_dir)
    text = None
    source = None

    for pdf in files["pdf"]:
        res = _pdf_page_words(pdf, page_idx)
        if res is not None:
            lines = _group_words_into_lines(res[0])
            text = "\n".join(ln["text"] for ln in lines)
            source = pdf
            break

    if text is None:
        words_list, _ = _load_words_list(files)
        if words_list is not None:
            pw = [{"text": w.get("text", ""), "bbox": (w.get("x0", 0), w.get("y0", 0), w.get("x1", 0), w.get("y1", 0))}
                  for w in words_list if int(w.get("page", 0) or 0) == page_idx + 1]
            text = "\n".join(ln["text"] for ln in _group_words_into_lines(pw))
            source = (files.get("words_json") or files.get("words_csv") or files.get("words_xlsx"))[0]

    if text is None and files["txt"]:
        try:
            with open(files["txt"][0], "r", encoding="utf-8") as f:
                _, sections = _parse_txt_sections(f.read())
            for num, body in sections:
                if num == page_idx + 1:
                    text = body
                    break
            if text is not None:
                source = files["txt"][0]
        except Exception:
            pass

    if text is None and source_path.lower().endswith(".pdf") and os.path.exists(source_path):
        res = _pdf_page_words(source_path, page_idx)
        if res is not None:
            text = "\n".join(ln["text"] for ln in _group_words_into_lines(res[0]))
            source = source_path

    text = text or ""
    chars = _count_chars(text)
    conf = estimate_text_layer_confidence(text.split(), text) if chars > 0 else 0.0
    # Prefer stored OCR confidence from s2pot_metadata.json
    meta_page = _get_meta_page(source_path, page_idx)
    if meta_page and meta_page.get("confidence") is not None and chars > 0:
        conf = float(meta_page["confidence"])

    editable = bool(files["pdf"] or files["docx"] or files["txt"] or files["hocr"]
                    or files["words_json"] or files["words_csv"] or files["words_xlsx"])
    return {
        "text": text,
        "chars": chars,
        "confidence": conf,
        "source": source,
        "has_ocr": chars > 0,
        "editable": editable,
        "manual_edit": bool(meta_page and meta_page.get("manual_edit")),
        "files": files
    }


def _get_meta_page(source_path, page_idx):
    meta = load_folder_metadata(os.path.dirname(source_path))
    if not meta:
        return None
    fmeta = meta.get("files", {}).get(os.path.basename(source_path))
    if not isinstance(fmeta, dict):
        return None
    pages = fmeta.get("pages") or []
    for p in pages:
        if isinstance(p, dict) and p.get("page_idx", -1) == page_idx:
            return p
    if page_idx < len(pages) and isinstance(pages[page_idx], dict):
        return pages[page_idx]
    return None


def get_document_pages_ocr_summary(source_path, n_pages, output_dir=None, data_dir=None):
    """
    Returns a list (one entry per page) of {"has_ocr", "chars", "confidence", "manual_edit"}.
    Uses s2pot_metadata.json when the document was processed, otherwise the result PDF
    text layer, otherwise the native text layer of the source PDF.
    """
    summary = [None] * n_pages
    meta = load_folder_metadata(os.path.dirname(source_path))
    fmeta = (meta or {}).get("files", {}).get(os.path.basename(source_path)) if meta else None
    if isinstance(fmeta, dict) and fmeta.get("status") == "done":
        for i, p in enumerate(fmeta.get("pages") or []):
            if not isinstance(p, dict):
                continue
            pidx = p.get("page_idx", i)
            if 0 <= pidx < n_pages and p.get("chars") is not None:
                c = int(p.get("chars") or 0)
                summary[pidx] = {
                    "has_ocr": c > 0,
                    "chars": c,
                    "confidence": p.get("confidence"),
                    "manual_edit": bool(p.get("manual_edit"))
                }

    if any(s is None for s in summary):
        files = find_result_files(source_path, output_dir, data_dir)
        candidates = list(files["pdf"])
        if source_path.lower().endswith(".pdf"):
            candidates.append(source_path)
        for pdf in candidates:
            try:
                with fitz.open(pdf) as doc:
                    for pidx in range(min(n_pages, len(doc))):
                        if summary[pidx] is not None:
                            continue
                        page = doc[pidx]
                        t = page.get_text("text").strip()
                        c = _count_chars(t)
                        conf = estimate_text_layer_confidence(page.get_text("words"), t) if c > 0 else 0.0
                        summary[pidx] = {"has_ocr": c > 0, "chars": c, "confidence": conf, "manual_edit": False}
            except Exception:
                continue
            if all(s is not None for s in summary):
                break

    return [s or {"has_ocr": False, "chars": 0, "confidence": None, "manual_edit": False} for s in summary]


# ----------------------------------------------------------------------------------------
# Layout of corrected text onto page coordinates
# ----------------------------------------------------------------------------------------
def _split_box_vertically(box, n):
    x0, y0, x1, y1 = box
    h = (y1 - y0) / float(max(1, n))
    return [(x0, y0 + i * h, x1, y0 + (i + 1) * h) for i in range(n)]


def _layout_new_lines(old_lines_norm, new_lines):
    """
    old_lines_norm: list of {"bbox": normalized (0..1) box, "text": str}
    new_lines: list of non-empty strings
    Returns list of (normalized_box, text) — one per new line.
    """
    if not new_lines:
        return []
    if not old_lines_norm:
        n = len(new_lines)
        h = min(0.025, 0.88 / float(n))
        return [((0.06, 0.06 + i * h * 1.2, 0.94, 0.06 + i * h * 1.2 + h), t) for i, t in enumerate(new_lines)]

    old_texts = [ln["text"] for ln in old_lines_norm]
    sm = difflib.SequenceMatcher(None, old_texts, new_lines, autojunk=False)
    placed = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal" or (tag == "replace" and (i2 - i1) == (j2 - j1)):
            for k in range(j2 - j1):
                placed.append((old_lines_norm[i1 + k]["bbox"], new_lines[j1 + k]))
        elif tag == "replace":
            boxes = [old_lines_norm[k]["bbox"] for k in range(i1, i2)]
            region = (min(b[0] for b in boxes), min(b[1] for b in boxes), max(b[2] for b in boxes), max(b[3] for b in boxes))
            for box, t in zip(_split_box_vertically(region, j2 - j1), new_lines[j1:j2]):
                placed.append((box, t))
        elif tag == "insert":
            anchor = old_lines_norm[i1 - 1]["bbox"] if i1 > 0 else old_lines_norm[0]["bbox"]
            ax0, ay0, ax1, ay1 = anchor
            h = max(0.005, ay1 - ay0)
            for k, t in enumerate(new_lines[j1:j2]):
                if i1 > 0:
                    ny0 = ay1 + k * h * 1.1 + h * 0.1
                else:
                    ny0 = ay0 - (j2 - j1 - k) * h * 1.1
                ny0 = max(0.0, min(1.0 - h, ny0))
                placed.append(((ax0, ny0, ax1, ny0 + h), t))
        # "delete": dropped lines are simply removed
    return placed


def _words_from_lines(placed):
    """Splits each placed line into word boxes proportional to character length (normalized coords)."""
    words = []
    for (x0, y0, x1, y1), line in placed:
        tokens = line.split()
        if not tokens:
            continue
        units = sum(len(t) for t in tokens) + (len(tokens) - 1)
        if units <= 0:
            continue
        unit_w = (x1 - x0) / float(units)
        cx = x0
        for t in tokens:
            wx1 = cx + len(t) * unit_w
            words.append({"text": t, "nbbox": (cx, y0, wx1, y1)})
            cx = wx1 + unit_w
    return words


def _reference_lines(source_path, page_idx, files):
    """
    Returns (old_lines_norm, px_size) where px_size = (W_px, H_px) of the word-coordinates
    pixel space (or None when unknown).
    """
    for pdf in files["pdf"]:
        res = _pdf_page_words(pdf, page_idx)
        if res is None:
            continue
        words, pw, ph, img_w, img_h = res
        lines = _group_words_into_lines(words)
        norm = [{"bbox": (l["bbox"][0] / pw, l["bbox"][1] / ph, l["bbox"][2] / pw, l["bbox"][3] / ph), "text": l["text"]} for l in lines]
        px = (img_w, img_h) if (img_w and img_h) else None
        return norm, px

    words_list, _ = _load_words_list(files)
    if words_list:
        W = max((float(w.get("x1", 0) or 0) for w in words_list), default=0) * 1.04 or 1.0
        H = max((float(w.get("y1", 0) or 0) for w in words_list), default=0) * 1.04 or 1.0
        pw = [{"text": w.get("text", ""), "bbox": (float(w.get("x0", 0)), float(w.get("y0", 0)), float(w.get("x1", 0)), float(w.get("y1", 0)))}
              for w in words_list if int(w.get("page", 0) or 0) == page_idx + 1]
        lines = _group_words_into_lines(pw)
        norm = [{"bbox": (l["bbox"][0] / W, l["bbox"][1] / H, l["bbox"][2] / W, l["bbox"][3] / H), "text": l["text"]} for l in lines]
        return norm, (W, H)

    return [], None


# ----------------------------------------------------------------------------------------
# Writers
# ----------------------------------------------------------------------------------------
def _safe_replace(tmp_path, final_path):
    try:
        os.replace(tmp_path, final_path)
    except PermissionError:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
        raise PermissionError(f"O ficheiro está aberto noutra aplicação: {os.path.basename(final_path)}")


def _update_pdf(pdf_path, page_idx, new_words_norm):
    doc = fitz.open(pdf_path)
    try:
        if page_idx >= len(doc):
            raise ValueError("página inexistente")
        page = doc[page_idx]
        # Remove the existing (invisible) text layer only — page images and vector graphics stay intact
        page.add_redact_annot(page.rect)
        kwargs = {"images": fitz.PDF_REDACT_IMAGE_NONE}
        if hasattr(fitz, "PDF_REDACT_LINE_ART_NONE"):
            kwargs["graphics"] = fitz.PDF_REDACT_LINE_ART_NONE
        page.apply_redactions(**kwargs)
        items = [{"text": w["text"], "bbox": w["nbbox"]} for w in new_words_norm]
        insert_invisible_text_items(page, items, scale_x=page.rect.width, scale_y=page.rect.height)
        tmp = pdf_path + ".s2pot_tmp"
        doc.save(tmp, garbage=3, deflate=True)
    finally:
        doc.close()
    _safe_replace(tmp, pdf_path)


def _update_txt(txt_path, page_idx, new_text):
    with open(txt_path, "r", encoding="utf-8") as f:
        content = f.read()
    preamble, sections = _parse_txt_sections(content)
    num = page_idx + 1
    if not sections:
        new_content = new_text
    else:
        for s in sections:
            if s[0] == num:
                s[1] = new_text
                break
        else:
            sections.append([num, new_text])
            sections.sort(key=lambda s: s[0])
        new_content = preamble + "\n\n".join(f"--- Page {n} ---\n{t}" for n, t in sections)
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(new_content)


def _update_hocr(hocr_path, page_idx, new_words_px):
    with open(hocr_path, "r", encoding="utf-8") as f:
        content = f.read()
    num = page_idx + 1
    starts = [m for m in re.finditer(r'<div class=["\']ocr_page["\'][^>]*>', content)]
    target = None
    for i, m in enumerate(starts):
        if re.search(rf'id=["\']page_{num}["\']', m.group(0)):
            target = (i, m)
            break
    if target is None:
        if len(starts) == 1 and page_idx == 0:
            target = (0, starts[0])
        else:
            raise ValueError(f"página {num} não encontrada no hOCR")
    i, m = target
    seg_start = m.start()
    seg_end = starts[i + 1].start() if i + 1 < len(starts) else len(content)
    segment = content[seg_start:seg_end]
    segment = re.sub(r'[ \t]*<span class=["\']ocrx_word["\'][^>]*>.*?</span>\n?', "", segment, flags=re.DOTALL)

    spans = "".join(
        f'     <span class="ocrx_word" id="word_{num}_{k + 1}" title="bbox {int(w["x0"])} {int(w["y0"])} {int(w["x1"])} {int(w["y1"])}">{html.escape(w["text"])}</span>\n'
        for k, w in enumerate(new_words_px)
    )
    par = re.search(r'<p class=["\']ocr_par["\'][^>]*>\n?', segment)
    if par:
        ins = par.end()
        if not segment[par.end() - 1:par.end()] == "\n":
            spans = "\n" + spans
    else:
        open_tag = re.match(r'<div class=["\']ocr_page["\'][^>]*>\n?', segment)
        ins = open_tag.end()
        if not segment[ins - 1:ins] == "\n":
            spans = "\n" + spans
    segment = segment[:ins] + spans + segment[ins:]
    with open(hocr_path, "w", encoding="utf-8") as f:
        f.write(content[:seg_start] + segment + content[seg_end:])


def _build_px_words(new_words_norm, px_size, base_name, page_num, template_words):
    W, H = px_size
    tmpl = template_words[0] if template_words else {}
    weight = tmpl.get("font_weight", "Normal")
    family = tmpl.get("font_family", "Sans-Serif")
    dpi_guess = 300.0
    out = []
    for w in new_words_norm:
        x0, y0, x1, y1 = w["nbbox"]
        px0, py0, px1, py1 = int(round(x0 * W)), int(round(y0 * H)), int(round(x1 * W)), int(round(y1 * H))
        h_px = max(1, py1 - py0)
        out.append({
            "filename": base_name,
            "page": page_num,
            "word_num": 0,
            "text": w["text"],
            "x0": px0, "y0": py0, "x1": px1, "y1": py1,
            "width": px1 - px0, "height": h_px,
            "font_size_pt": round(h_px * 72.0 / dpi_guess, 1),
            "font_weight": weight,
            "font_family": family
        })
    return out


def _merge_page_words(words_list, page_num, new_page_words):
    before = [w for w in words_list if int(w.get("page", 0) or 0) < page_num]
    after = [w for w in words_list if int(w.get("page", 0) or 0) > page_num]
    merged = before + new_page_words + after
    for i, w in enumerate(merged):
        w["word_num"] = i + 1
    return merged


def _update_docx(docx_path, page_idx, new_lines):
    if docx is None:
        raise RuntimeError("python-docx indisponível")
    d = docx.Document(docx_path)
    body = d.element.body
    W_P, W_TBL = qn("w:p"), qn("w:tbl")

    def is_page_break(el):
        return el.tag == W_P and any(br.get(qn("w:type")) == "page" for br in el.iter(qn("w:br")))

    def has_drawing(el):
        return any(True for _ in el.iter(qn("w:drawing"))) or any(True for _ in el.iter(qn("w:pict")))

    def el_text(el):
        return "".join(t.text or "" for t in el.iter(qn("w:t")))

    # Split body children into pages (page-break paragraphs act as separators)
    pages = [[]]
    for el in list(body):
        if el.tag not in (W_P, W_TBL):
            continue
        if is_page_break(el):
            pages.append([])
            pages[-1].append(("break", el))
            continue
        pages[-1].append(("el", el))
    if page_idx >= len(pages):
        raise ValueError("página inexistente no DOCX")

    seg = [el for kind, el in pages[page_idx] if kind == "el"]
    text_paras = [el for el in seg if el.tag == W_P and not has_drawing(el) and el_text(el).strip()]

    template_p = text_paras[0] if text_paras else None
    template_r = None
    if template_p is not None:
        runs = list(template_p.iter(qn("w:r")))
        template_r = runs[0] if runs else None

    # Insertion anchor
    if text_paras:
        anchor = text_paras[0]
        insert_before = True
    else:
        nxt = pages[page_idx + 1][0][1] if page_idx + 1 < len(pages) and pages[page_idx + 1] else None
        if nxt is not None:
            anchor, insert_before = nxt, True
        elif seg:
            anchor, insert_before = seg[-1], False
        else:
            anchor, insert_before = None, False

    new_elements = []
    for line in new_lines:
        if template_p is not None:
            p_el = copy.deepcopy(template_p)
            for child in list(p_el):
                if child.tag != qn("w:pPr"):
                    p_el.remove(child)
            r_el = copy.deepcopy(template_r) if template_r is not None else None
            if r_el is None:
                tmp_par = d.add_paragraph()
                tmp_par.add_run(line)
                r_el = list(tmp_par._p.iter(qn("w:r")))[0]
                body.remove(tmp_par._p)
            else:
                for child in list(r_el):
                    if child.tag != qn("w:rPr"):
                        r_el.remove(child)
                t_el = r_el.makeelement(qn("w:t"), {})
                t_el.text = line
                t_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                r_el.append(t_el)
            p_el.append(r_el)
        else:
            par = d.add_paragraph()
            par.paragraph_format.space_before = Pt(0)
            par.paragraph_format.space_after = Pt(1.0)
            run = par.add_run(line)
            run.font.size = Pt(9.5)
            run.font.name = "Calibri"
            p_el = par._p
            body.remove(p_el)
        new_elements.append(p_el)

    if anchor is None:
        sect = body.find(qn("w:sectPr"))
        for p_el in new_elements:
            if sect is not None:
                sect.addprevious(p_el)
            else:
                body.append(p_el)
    elif insert_before:
        for p_el in new_elements:
            anchor.addprevious(p_el)
    else:
        for p_el in reversed(new_elements):
            anchor.addnext(p_el)

    for el in text_paras:
        el.getparent().remove(el)

    tmp = docx_path + ".s2pot_tmp"
    d.save(tmp)
    _safe_replace(tmp, docx_path)


def _update_metadata(source_path, page_idx, chars, output_dirs=()):
    """Updates s2pot_metadata.json in the input folder. Returns the updated file metadata dict or None."""
    folder = os.path.dirname(source_path)
    fname = os.path.basename(source_path)
    meta = load_folder_metadata(folder)
    if not meta:
        return None
    fmeta = meta.get("files", {}).get(fname)
    if not isinstance(fmeta, dict):
        return None
    pages = fmeta.setdefault("pages", [])
    target = None
    for i, p in enumerate(pages):
        if isinstance(p, dict) and p.get("page_idx", i) == page_idx:
            target = p
            break
    if target is None:
        while len(pages) <= page_idx:
            pages.append({"page_idx": len(pages)})
        target = pages[page_idx]
    target["chars"] = chars
    target["confidence"] = MANUAL_EDIT_CONFIDENCE if chars > 0 else 0.0
    target["manual_edit"] = True
    target["edited_at"] = datetime.now().isoformat(timespec="seconds")

    pl = [p for p in pages if isinstance(p, dict)]
    tot = sum(int(p.get("chars") or 0) for p in pl)
    fmeta["total_chars"] = tot
    conf_pages = [p for p in pl if p.get("confidence") is not None]
    if conf_pages:
        if tot > 0:
            fmeta["confidence"] = round(sum((p.get("confidence") or 0.0) * int(p.get("chars") or 0) for p in conf_pages) / float(tot), 1)
        else:
            fmeta["confidence"] = round(sum((p.get("confidence") or 0.0) for p in conf_pages) / float(len(conf_pages)), 1)
        fmeta["min_page_confidence"] = round(min((p.get("confidence") or 0.0) for p in conf_pages), 1)
        fmeta["low_conf_pages"] = sum(1 for p in conf_pages if ((p.get("confidence") or 0.0) < 80.0 or int(p.get("chars") or 0) == 0))
    meta["updated_at"] = datetime.now().isoformat()
    save_folder_metadata(folder, meta)
    return fmeta


def save_edited_page_ocr_text(source_path, page_idx, new_text, output_dir=None, data_dir=None):
    """
    Propagates the corrected OCR text of one page to every generated result file.
    Returns {"updated": [paths], "errors": [(path, msg)], "chars": int, "file_meta": dict|None}.
    """
    files = find_result_files(source_path, output_dir, data_dir)
    base = files["base_name"]
    page_num = page_idx + 1

    new_text = (new_text or "").replace("\r\n", "\n").replace("\r", "\n")
    new_lines = [ln.strip() for ln in new_text.split("\n") if ln.strip()]
    clean_text = "\n".join(new_lines)
    chars = _count_chars(clean_text)

    old_lines_norm, px_size = _reference_lines(source_path, page_idx, files)
    placed = _layout_new_lines(old_lines_norm, new_lines)
    new_words_norm = _words_from_lines(placed)

    updated, errors = [], []

    for pdf in files["pdf"]:
        try:
            _update_pdf(pdf, page_idx, new_words_norm)
            updated.append(pdf)
        except Exception as e:
            errors.append((pdf, str(e)))

    for docx_path in files["docx"]:
        try:
            _update_docx(docx_path, page_idx, new_lines)
            updated.append(docx_path)
        except Exception as e:
            errors.append((docx_path, str(e)))

    for txt in files["txt"]:
        try:
            _update_txt(txt, page_idx, clean_text)
            updated.append(txt)
        except Exception as e:
            errors.append((txt, str(e)))

    words_list, json_payload = _load_words_list(files)
    if px_size is None and words_list:
        W = max((float(w.get("x1", 0) or 0) for w in words_list), default=0) * 1.04 or 2480.0
        H = max((float(w.get("y1", 0) or 0) for w in words_list), default=0) * 1.04 or 3508.0
        px_size = (W, H)
    if px_size is None:
        px_size = (2480.0, 3508.0)
    template_words = [w for w in (words_list or []) if int(w.get("page", 0) or 0) == page_num]
    new_px_words = _build_px_words(new_words_norm, px_size, base, page_num, template_words)

    for hocr in files["hocr"]:
        try:
            _update_hocr(hocr, page_idx, new_px_words)
            updated.append(hocr)
        except Exception as e:
            errors.append((hocr, str(e)))

    if words_list is not None:
        merged = _merge_page_words(words_list, page_num, [dict(w) for w in new_px_words])
        total_pages = (json_payload or {}).get("total_pages") or max((int(w.get("page", 0) or 0) for w in merged), default=page_num)
        for path in files["words_json"]:
            try:
                write_words_json(path, (json_payload or {}).get("filename", base), merged, total_pages)
                updated.append(path)
            except Exception as e:
                errors.append((path, str(e)))
        for path in files["words_csv"]:
            try:
                write_words_csv(path, merged)
                updated.append(path)
            except Exception as e:
                errors.append((path, str(e)))
        for path in files["words_xlsx"]:
            try:
                write_words_xlsx(path, merged)
                updated.append(path)
            except Exception as e:
                errors.append((path, str(e)))

    file_meta = None
    try:
        file_meta = _update_metadata(source_path, page_idx, chars)
    except Exception as e:
        errors.append(("s2pot_metadata.json", str(e)))

    return {"updated": updated, "errors": errors, "chars": chars, "file_meta": file_meta}
